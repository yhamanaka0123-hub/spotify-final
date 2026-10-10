"""Golden-50 evaluation: human labels (Excel sheet) vs pipeline predictions. No model calls.

    .venv/bin/python -m evals.golden --xlsx "/mnt/c/Users/yhama/Downloads/Final assignment/golden50_labeling_ja.xlsx" \
        --db work/full.db

Human labels never enter prompts: this script only reads them after classification.
Writes evals/golden_labels.csv (human), evals/golden_report.json (metrics, confusion, disagreements).
"""
import argparse
import csv
import json
import sqlite3
from collections import Counter
from pathlib import Path

from openpyxl import load_workbook

from pipeline.contract import INTENTS, TOPICS
from pipeline.enrich import LABEL_CONFIG

HERE = Path(__file__).resolve().parent
SENTIMENT_TOLERANCE = 0.5  # predeclared: |pred - human| <= 0.5 counts as agreement


def code(v):
    return None if v in (None, "") else str(v).split(" | ")[0].strip()


def read_human(xlsx):
    ws = load_workbook(xlsx, data_only=True)["ラベル付け"]
    head = [c.value for c in ws[1]]
    col = {name: head.index(name) for name in head}
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if not r[col["review_id"]]:
            continue
        rows.append({
            "review_id": r[col["review_id"]], "review_text": r[col["原文 (review_text)"]],
            "topic": code(r[col["topic"]]), "intent": code(r[col["intent"]]),
            "severity": int(code(r[col["severity"]])) if code(r[col["severity"]]) else None,
            "sentiment": float(code(r[col["sentiment"]])) if code(r[col["sentiment"]]) else None,
            "evidence_quote": r[col["evidence_quote（原文からそのままコピー）"]] or "",
            "entities": r[col["entities（任意・カンマ区切り）"]] or "",
            "needs_review": code(r[col["needs_review"]]) == "true",
            "note": r[col["メモ（任意）"]] or "",
        })
    return rows


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--xlsx", required=True)
    ap.add_argument("--db", required=True)
    args = ap.parse_args()
    human = read_human(args.xlsx)
    missing = [h["review_id"] for h in human if not (h["topic"] and h["intent"] and h["severity"])]
    con = sqlite3.connect(args.db)
    pred = {rid: dict(zip(("topic", "intent", "severity", "sentiment", "evidence_quote", "needs_review", "decided_by"), rest))
            for rid, *rest in con.execute("""
        SELECT rv.review_id, r.topic, r.intent, r.severity, r.sentiment, r.evidence_quote, r.needs_review, r.decided_by
        FROM reviews rv JOIN results r ON r.text_key = rv.text_key AND r.label_config = ?""", (LABEL_CONFIG,))}

    with open(HERE / "golden_labels.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(human[0]))
        w.writeheader()
        w.writerows(human)

    cases, n = [], 0
    agree = Counter()
    sev_abs, sent_abs = [], []
    conf_topic, conf_intent = Counter(), Counter()
    for h in human:
        p = pred.get(h["review_id"])
        case = {"review_id": h["review_id"], "text": h["review_text"], "human": {k: h[k] for k in
                ("topic", "intent", "severity", "sentiment", "needs_review")}, "pred": p, "note": h["note"]}
        if p is None or h["review_id"] in missing:
            case["status"] = "missing_prediction" if p is None else "missing_human_label"
            cases.append(case)
            continue
        n += 1
        for f in ("topic", "intent", "severity"):
            agree[f] += p[f] == h[f]
        sev_abs.append(abs(p["severity"] - h["severity"]))
        if h["sentiment"] is not None:
            d = abs(p["sentiment"] - h["sentiment"])
            sent_abs.append(d)
            agree["sentiment_within_tol"] += d <= SENTIMENT_TOLERANCE
        agree["quote_is_substring"] += p["evidence_quote"] in h["review_text"]
        agree["needs_review"] += bool(p["needs_review"]) == h["needs_review"]
        conf_topic[(h["topic"], p["topic"])] += 1
        conf_intent[(h["intent"], p["intent"])] += 1
        case["disagree"] = [f for f in ("topic", "intent", "severity") if p[f] != h[f]]
        cases.append(case)

    def macro_f1(conf, labels):
        f1s = []
        for lab in labels:
            tp = conf[(lab, lab)]
            fp = sum(v for (h, p), v in conf.items() if p == lab and h != lab)
            fn = sum(v for (h, p), v in conf.items() if h == lab and p != lab)
            if tp + fp + fn:
                f1s.append(2 * tp / (2 * tp + fp + fn))
        return round(sum(f1s) / len(f1s), 4) if f1s else None

    report = {
        "label_config": LABEL_CONFIG, "evaluated": n, "missing": len(human) - n,
        "agreement": {"topic": round(agree["topic"] / n, 4), "intent": round(agree["intent"] / n, 4),
                      "severity_exact": round(agree["severity"] / n, 4),
                      "severity_mae": round(sum(sev_abs) / n, 4),
                      "sentiment_mae": round(sum(sent_abs) / len(sent_abs), 4) if sent_abs else None,
                      f"sentiment_within_{SENTIMENT_TOLERANCE}": round(agree["sentiment_within_tol"] / max(1, len(sent_abs)), 4),
                      "quote_is_exact_substring": round(agree["quote_is_substring"] / n, 4),
                      "needs_review_match": round(agree["needs_review"] / n, 4),
                      "topic_macro_f1": macro_f1(conf_topic, TOPICS), "intent_macro_f1": macro_f1(conf_intent, INTENTS)},
        "human_flagged_ambiguous": sum(h["needs_review"] for h in human),
        "per_topic_counts": {"human": Counter(h["topic"] for h in human), "pred": Counter(
            pred[h["review_id"]]["topic"] for h in human if h["review_id"] in pred)},
        "confusion_topic": {f"{h}->{p}": v for (h, p), v in sorted(conf_topic.items()) if h != p},
        "confusion_intent": {f"{h}->{p}": v for (h, p), v in sorted(conf_intent.items()) if h != p},
        "cases": cases,
    }
    (HERE / "golden_report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2, default=dict) + "\n",
                                             encoding="utf-8")
    print(json.dumps({k: report[k] for k in ("evaluated", "missing", "agreement", "human_flagged_ambiguous")}, indent=2))


if __name__ == "__main__":
    main()
