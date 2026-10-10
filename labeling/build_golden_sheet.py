"""Build a Japanese-annotated Excel sheet for hand-labeling the golden 50.

Usage:
    .venv/bin/python labeling/build_golden_sheet.py

Reads data/golden_50_to_label.csv and labeling/golden50_translations_ja.json,
writes labeling/golden50_labeling_ja.xlsx. No labels are pre-filled: the
golden set must be human-labeled (see GRADING_CONTRACT.md).
"""
import csv
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "data" / "golden_50_to_label.csv"
TRANSLATIONS = ROOT / "labeling" / "golden50_translations_ja.json"
OUT = ROOT / "labeling" / "golden50_labeling_ja.xlsx"

# Dropdown values: English code first (used by the converter), Japanese hint after " | ".
TOPICS = [
    ("access", "ログイン・登録・パスワード・アカウント"),
    ("usability", "操作性・画面・キュー/プレイリスト管理・広告による中断"),
    ("playback", "再生できない・クラッシュ・遅い・接続・音質"),
    ("downloads", "ダウンロード・オフライン再生・保存曲が消える"),
    ("catalog", "曲がない・検索・おすすめ・歌詞"),
    ("billing", "料金・請求・サブスク・プレミアム限定機能"),
    ("support", "サポートへの問い合わせと対応"),
    ("other", "一般的な称賛/批判・無関係・特定トピックなし"),
]
INTENTS = [
    ("cancellation", "やめる・アンインストール・解約（予告含む）"),
    ("complaint", "不満・悪い体験（褒め＋批判の混在も）"),
    ("request", "要望（不具合の報告なし）"),
    ("praise", "称賛"),
    ("unclear", "意味不明・無関係・ボイコットの標語だけ"),
]
SEVERITIES = [
    ("1", "問題なし（称賛・不明・純粋な要望）"),
    ("2", "不満・一般的な批判・軽い迷惑（機能は失われていない）"),
    ("3", "機能が低下/制限されているが一部は使える"),
    ("4", "主要な作業が完全にできない（ログイン不可・再生不可など）"),
    ("5", "金銭・プライバシー・データの深刻な被害が明記"),
]
SENTIMENTS = [("-1", "とても否定的"), ("-0.5", "やや否定的"), ("0", "中立"),
              ("0.5", "やや肯定的"), ("1", "とても肯定的")]
NEEDS_REVIEW = [("false", "迷いなし"), ("true", "判断に迷う・文脈不足")]


def opt(code, hint):
    return f"{code} | {hint}"


def add_list(ws, sheet_title, col, options, rng):
    """Put options on the lists sheet and attach a dropdown referencing them."""
    for i, (code, hint) in enumerate(options, 1):
        ws.cell(row=i, column=col, value=opt(code, hint))
    letter = ws.cell(row=1, column=col).column_letter
    dv = DataValidation(type="list", formula1=f"='{ws.title}'!${letter}$1:${letter}${len(options)}",
                        allow_blank=True, showDropDown=False)
    dv.error = "リストから選んでください"
    dv.add(rng)
    return dv


def main():
    rows = list(csv.DictReader(SRC.open(encoding="utf-8", newline="")))
    ja = json.loads(TRANSLATIONS.read_text(encoding="utf-8"))
    assert len(rows) == len(ja) == 50, (len(rows), len(ja))

    wb = Workbook()
    ws = wb.active
    ws.title = "ラベル付け"
    lists = wb.create_sheet("選択肢")

    headers = ["No", "★", "原文 (review_text)", "日本語訳（参考）",
               "topic", "intent", "severity", "sentiment",
               "evidence_quote（原文からそのままコピー）", "entities（任意・カンマ区切り）",
               "needs_review", "メモ（任意）", "review_id"]
    ws.append(headers)
    for i, (r, t) in enumerate(zip(rows, ja), 1):
        ws.append([i, int(r["review_rating"]), r["review_text"], t,
                   None, None, None, None, None, None, None, None, r["review_id"]])

    n = len(rows) + 1
    sheet = ws.title
    for dv in (
        add_list(lists, sheet, 1, TOPICS, f"E2:E{n}"),
        add_list(lists, sheet, 2, INTENTS, f"F2:F{n}"),
        add_list(lists, sheet, 3, SEVERITIES, f"G2:G{n}"),
        add_list(lists, sheet, 4, SENTIMENTS, f"H2:H{n}"),
        add_list(lists, sheet, 5, NEEDS_REVIEW, f"K2:K{n}"),
    ):
        ws.add_data_validation(dv)
    lists.sheet_state = "hidden"

    widths = [5, 4, 50, 50, 26, 22, 22, 16, 40, 22, 16, 24, 38]
    for col, w in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = w
    head_fill = PatternFill("solid", fgColor="1F4E78")
    input_fill = PatternFill("solid", fgColor="FFF2CC")
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = head_fill
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for row in ws.iter_rows(min_row=2, max_row=n):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
        for c in row[4:11]:
            c.fill = input_fill
    ws.freeze_panes = "E2"

    guide = wb.create_sheet("早見表", 0)
    lines = [
        ("Golden 50 ラベル付け 早見表", None),
        ("", None),
        ("やり方", None),
        ("1. 「ラベル付け」シートの黄色い列を埋める（topic / intent / severity / sentiment / needs_review はプルダウン）", None),
        ("2. evidence_quote は【原文】から根拠部分をそのままコピペ（日本語訳ではなく英語などの原文。1文字でも変えない）", None),
        ("3. 日本語訳は参考。判断は必ず原文で。星（★）は判断に使わない（本文だけで判断）", None),
        ("4. 迷ったら needs_review = true にして、メモに理由を書く", None),
        ("", None),
        ("topic（何についてか）", None),
        *[(opt(c, h), None) for c, h in TOPICS],
        ("・複数の問題があれば一番深刻なもの。同じなら最初に出てきたもの", None),
        ("・称賛なら最初に褒めている具体的機能。漠然とした称賛は other", None),
        ("・「プレミアム」と書いてあるだけでは billing ではない（課金者の再生クラッシュは playback）", None),
        ("", None),
        ("intent（優先順：上にあるものが勝つ）", None),
        *[(opt(c, h), None) for c, h in INTENTS],
        ("・「ダメなアプリ」は complaint（具体的な不具合を推測しない）", None),
        ("・ボイコットの標語や無関係な文は unclear（製品への不満か、本人がやめると明言していれば別）", None),
        ("", None),
        ("severity（深刻度）", None),
        *[(opt(c, h), None) for c, h in SEVERITIES],
        ("・解約の意思だけで深刻度は上げない。高い料金・クラッシュ・怒った口調だけでは 5 にしない", None),
        ("", None),
        ("sentiment（感情）: -1 〜 1", None),
        *[(opt(c, h), None) for c, h in SENTIMENTS],
        ("", None),
        ("entities（任意）: 本文に出てくる機能名など（例: ads, shuffle, lyrics, offline mode）。なければ空欄", None),
    ]
    for text, _ in lines:
        guide.append([text])
    guide.column_dimensions["A"].width = 120
    for c in guide["A"]:
        if c.value and not c.value.startswith(("・", "1.", "2.", "3.", "4.")) and " | " not in c.value:
            c.font = Font(bold=True, size=14 if c.row == 1 else 11)

    wb.active = 1
    wb.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    main()
